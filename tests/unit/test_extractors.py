from pathlib import Path

import docx
import pytest

from cvopt.adapters.extractors import LocalTextExtractor
from cvopt.adapters.language import LinguaDetector
from cvopt.config import DocumentSettings
from cvopt.errors import InputError
from tests.scenarios import CV, LETTER

EXTRACT = LocalTextExtractor(DocumentSettings())


async def test_pdf_text_layer() -> None:
    out = await EXTRACT.extract(CV)
    assert out.media_type == "application/pdf" and not out.needs_transcription
    assert "Nordlicht Analytics AG" in out.text and out.pages == 1


async def test_short_pdf_needs_transcription() -> None:
    out = await EXTRACT.extract(Path(__file__).parents[1] / "golden" / "note.pdf")
    assert out.needs_transcription and out.text == ""


async def test_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_paragraph("First paragraph")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Python", "5 years"
    d.add_paragraph("After the table")
    path = tmp_path / "cv.docx"
    d.save(str(path))
    out = await EXTRACT.extract(path)
    assert out.text.splitlines() == ["First paragraph", "Python | 5 years", "After the table"]


async def test_text_images_and_refusals(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_bytes("Zürich".encode("cp1252"))
    assert (await EXTRACT.extract(tmp_path / "a.txt")).text == "Zürich"
    (tmp_path / "a.png").write_bytes(b"\x89PNG")
    assert (await EXTRACT.extract(tmp_path / "a.png")).needs_transcription
    (tmp_path / "a.xls").write_bytes(b"x")
    with pytest.raises(InputError):
        await EXTRACT.extract(tmp_path / "a.xls")
    (tmp_path / "fake.pdf").write_bytes(b"not a pdf")
    with pytest.raises(InputError):
        await EXTRACT.extract(tmp_path / "fake.pdf")
    with pytest.raises(InputError):
        await LocalTextExtractor(DocumentSettings(max_mb=0.00001)).extract(CV)
    with pytest.raises(InputError):
        await EXTRACT.extract(tmp_path / "missing.pdf")


async def test_language_detection() -> None:
    detector = LinguaDetector()
    assert detector.detect((await EXTRACT.extract(CV)).text) == "en"
    assert detector.detect((await EXTRACT.extract(LETTER)).text) == "de"
    assert detector.detect("") == "und"
